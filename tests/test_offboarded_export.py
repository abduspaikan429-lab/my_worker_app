import pytest
from pathlib import Path
import openpyxl
from scripts.export_offboarded_workers import (
    COLUMNS,
    TEXT_FORMAT_COLUMNS,
    build_excel_workbook,
    map_row,
)
from modules.master_data import load_master_df
from modules.report_generator import extract_offboarded_report


def test_columns_count_and_order():
    expected_cols = [
        "工人姓名", "性别", "出生日期", "年龄", "籍贯", "民族", "政治面貌", "工种",
        "家庭住址", "身份证号码", "发证机关", "有效起始日期", "有效截止日期",
        "开户银行", "开户支行", "工资卡号", "单位工资", "劳动合同登记",
        "劳动合同编号", "安全培训情况", "进场教育", "首次进场时间", "进场时间",
        "退场时间", "退场工资结清情况", "持上岗证情况", "联系电话",
    ]
    assert COLUMNS == expected_cols
    assert len(COLUMNS) == 27


def test_build_excel_workbook_and_verify_sheets(tmp_path):
    target_excel = tmp_path / "test_offboarded.xlsx"
    build_excel_workbook(target_excel)
    assert target_excel.exists()

    wb = openpyxl.load_workbook(target_excel)
    assert "退场人员名单" in wb.sheetnames
    assert "含历史退场记录人员(37人)" in wb.sheetnames

    ws1 = wb["退场人员名单"]
    headers = [cell.value for cell in ws1[1]]
    assert headers == COLUMNS
    assert ws1.max_row >= 25  # 1 header + offboarded workers

    # 验证纯文本格式 @
    col_idx_id = COLUMNS.index("身份证号码") + 1
    col_idx_card = COLUMNS.index("工资卡号") + 1
    cell_id = ws1.cell(row=2, column=col_idx_id)
    cell_card = ws1.cell(row=2, column=col_idx_card)
    assert cell_id.number_format == "@"
    assert cell_card.number_format == "@"


def test_extract_offboarded_report_dataframe():
    df = load_master_df()
    off_df = df[df["在场/进退场状态"].astype(str).str.strip().isin(["已退场", "离场"])].copy()
    assert len(off_df) >= 20

    res = extract_offboarded_report(off_df, ensure_text=True)
    assert list(res.columns) == COLUMNS
    assert len(res) == len(off_df)

    # 验证防科学计数法单引号
    first_id = res.iloc[0]["身份证号码"]
    assert first_id.startswith("'")
