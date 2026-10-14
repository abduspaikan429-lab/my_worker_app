# tests/test_payroll_worker_sort.py
import unittest
import pandas as pd
from services.payroll_calculation_service import (
    parse_worker_sort_order,
    sort_workers_df_by_order
)

class TestPayrollWorkerSort(unittest.TestCase):
    def setUp(self):
        self.sample_df = pd.DataFrame([
            {'姓名': '张三', '工种': '木工', '最终核定天数': 25.0, '应发工资': 7500.0},
            {'姓名': '李四', '工种': '钢筋工', '最终核定天数': 20.0, '应发工资': 6000.0},
            {'姓名': '王五', '工种': '瓦工', '最终核定天数': 18.0, '应发工资': 5400.0},
            {'姓名': '赵六', '工种': '油漆工', '最终核定天数': 15.0, '应发工资': 4500.0},
            {'姓名': '孙七', '工种': '架子工', '最终核定天数': 10.0, '应发工资': 3000.0},
        ])

    def test_parse_worker_sort_order_plain(self):
        text = "王五\n张三\n李四"
        names = parse_worker_sort_order(text, self.sample_df['姓名'].tolist())
        self.assertEqual(names, ['王五', '张三', '李四'])

    def test_parse_worker_sort_order_with_numbers_and_punctuation(self):
        text = """1. 王五
2、 李四
3 赵六
(4) 孙七
"""
        names = parse_worker_sort_order(text, self.sample_df['姓名'].tolist())
        self.assertEqual(names, ['王五', '李四', '赵六', '孙七'])

    def test_parse_worker_sort_order_from_excel_copy(self):
        text = """1\t王五\t瓦工
2\t张三\t木工
"""
        names = parse_worker_sort_order(text, self.sample_df['姓名'].tolist())
        self.assertEqual(names, ['王五', '张三'])

    def test_parse_worker_sort_order_comma_separated(self):
        text = "王五, 李四，赵六、孙七"
        names = parse_worker_sort_order(text, self.sample_df['姓名'].tolist())
        self.assertEqual(names, ['王五', '李四', '赵六', '孙七'])

    def test_sort_workers_df_full_order(self):
        custom_order = ['孙七', '赵六', '王五', '李四', '张三']
        sorted_df, stats = sort_workers_df_by_order(self.sample_df, custom_order)
        self.assertEqual(sorted_df['姓名'].tolist(), custom_order)
        self.assertEqual(stats['matched'], custom_order)
        self.assertEqual(stats['unmatched'], [])
        self.assertEqual(stats['appended'], [])

    def test_sort_workers_df_partial_order(self):
        # 用户只指定了前两人，剩下的三人应自动保持原顺序追加在末尾
        custom_order = ['王五', '张三']
        sorted_df, stats = sort_workers_df_by_order(self.sample_df, custom_order)
        self.assertEqual(sorted_df['姓名'].tolist(), ['王五', '张三', '李四', '赵六', '孙七'])
        self.assertEqual(stats['matched'], ['王五', '张三'])
        self.assertEqual(stats['appended'], ['李四', '赵六', '孙七'])

    def test_sort_workers_df_with_extra_names(self):
        # 包含了不在名单里的“外人”
        custom_order = ['王五', '外人', '李四']
        sorted_df, stats = sort_workers_df_by_order(self.sample_df, custom_order)
        self.assertEqual(sorted_df['姓名'].tolist(), ['王五', '李四', '张三', '赵六', '孙七'])
        self.assertEqual(stats['matched'], ['王五', '李四'])
        self.assertEqual(stats['unmatched'], ['外人'])
        self.assertEqual(stats['appended'], ['张三', '赵六', '孙七'])

    def test_sort_workers_df_empty_order(self):
        sorted_df, stats = sort_workers_df_by_order(self.sample_df, [])
        self.assertEqual(sorted_df['姓名'].tolist(), self.sample_df['姓名'].tolist())
        self.assertEqual(stats['matched'], [])
        self.assertEqual(stats['appended'], self.sample_df['姓名'].tolist())

    def test_excel_export_preserves_custom_order(self):
        import io
        import zipfile
        from openpyxl import load_workbook
        from modules.attendance_payroll import build_all_exports_zip

        test_df = pd.DataFrame([
            {'分包/所属企业': '江苏旭之升建筑工程有限公司', '月份': '8月', '姓名': '张三', '工种': '木工', '最终核定天数': 20.0, '应发工资': 6000.0, '日薪': 300.0, '银行卡号': '62220001', '提取日薪': 300.0},
            {'分包/所属企业': '江苏旭之升建筑工程有限公司', '月份': '8月', '姓名': '李四', '工种': '电工', '最终核定天数': 22.0, '应发工资': 6600.0, '日薪': 300.0, '银行卡号': '62220002', '提取日薪': 300.0},
            {'分包/所属企业': '江苏旭之升建筑工程有限公司', '月份': '8月', '姓名': '王五', '工种': '瓦工', '最终核定天数': 18.0, '应发工资': 5400.0, '日薪': 300.0, '银行卡号': '62220003', '提取日薪': 300.0},
        ])

        # 自定义顺序：王五 -> 张三 -> 李四
        target_order = ['王五', '张三', '李四']
        sorted_df, _ = sort_workers_df_by_order(test_df, target_order)
        zip_bytes = build_all_exports_zip(sorted_df)
        self.assertTrue(len(zip_bytes) > 0)

        zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
        excel_name = [n for n in zf.namelist() if n.endswith('.xlsx')][0]
        wb = load_workbook(io.BytesIO(zf.read(excel_name)))

        att_ws = wb['考勤表']
        wage_ws = wb['工资确认表']

        att_names = [att_ws.cell(row=r, column=2).value for r in range(6, 6 + 3)]
        wage_names = [wage_ws.cell(row=r, column=2).value for r in range(5, 5 + 3)]

        self.assertEqual(att_names, target_order)
        self.assertEqual(wage_names, target_order)


if __name__ == '__main__':
    unittest.main()

