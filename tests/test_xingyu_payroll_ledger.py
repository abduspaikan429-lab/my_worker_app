# -*- coding: utf-8 -*-
"""
tests/test_xingyu_payroll_ledger.py
测试上海兴豫建筑工程有限公司标准工资台账生成服务
"""
import os
import openpyxl
from services.xingyu_payroll_service import (
    build_xingyu_payroll_workbook,
    export_xingyu_excel,
    get_xingyu_excel_bytes
)


def test_build_xingyu_workbook():
    wb = build_xingyu_payroll_workbook()
    ws = wb.active
    
    assert ws.title == "工资台账"
    # 78 列
    assert ws.max_column == 78
    # 82 行 (大标题1 + 表头2 + 数据78 + 合计1)
    assert ws.max_row == 82
    # 冻结窗格
    assert ws.freeze_panes == "F4"
    # 标题
    assert "上海兴豫建筑工程有限公司工资台账" in ws["A1"].value
    
    # 第4行（杨成）
    assert ws["A4"].value == 1
    assert ws["B4"].value == "杨成"
    assert ws["E4"].value == "2025-02-19"
    assert ws["BS4"].value == "=SUM(F4:J4)"
    assert ws["BT4"].value == "=SUM(K4:O4)"
    assert ws["BV4"].value == "=BS4-BT4"
    
    # 合计行
    assert ws["A82"].value == "合  计"
    assert ws["BS82"].value == "=SUM(BS4:BS81)"


def test_export_xingyu_excel(tmp_path):
    target = os.path.join(tmp_path, "test_output.xlsx")
    res_path = export_xingyu_excel(target)
    assert os.path.exists(res_path)
    
    # 检查能否正常读取
    wb = openpyxl.load_workbook(res_path)
    assert "工资台账" in wb.sheetnames


def test_get_xingyu_excel_bytes():
    b = get_xingyu_excel_bytes()
    assert len(b) > 5000
    assert b[:4] == b"PK\x03\x04"  # zip/xlsx 魔数
